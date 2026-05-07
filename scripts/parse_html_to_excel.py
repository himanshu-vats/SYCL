"""
Parse the saved CricClubs batting HTML page → Excel.
Usage: python3 parse_html_to_excel.py batting_page.html
(Save the page as HTML from your browser: Ctrl+S or right-click > Save As)
"""
import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from bs4 import BeautifulSoup

def parse(html_path, out_path="batting_records.xlsx"):
    with open(html_path, encoding="utf-8", errors="ignore") as f:
        soup = BeautifulSoup(f.read(), "html.parser")

    table = soup.find("table", id="webrecordtable") or soup.find("table", class_="playersData")
    if not table:
        print("ERROR: Could not find the data table in the HTML file.")
        sys.exit(1)

    rows = []
    for tr in table.find("tbody").find_all("tr"):
        tds = tr.find_all("td")
        if len(tds) < 14:
            continue
        rows.append([td.get_text(strip=True) for td in tds])

    headers = ["#", "Player", "Team", "Mat", "Inns", "NO", "Runs",
               "4s", "6s", "50s", "100s", "HS", "SR", "Avg", "Points"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Batting Records"

    hdr_fill = PatternFill("solid", fgColor="0F2A1A")
    hdr_font = Font(bold=True, color="FFFFFF", size=11)
    alt_fill = PatternFill("solid", fgColor="E8F5E9")

    ws.append(headers)
    for cell in ws[1]:
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center")

    for i, row in enumerate(rows, 2):
        # Pad or trim to match header count
        row_data = row[:len(headers)]
        while len(row_data) < len(headers):
            row_data.append("")
        ws.append(row_data)
        if i % 2 == 0:
            for cell in ws[i]:
                cell.fill = alt_fill
        ws.cell(i, 2).font = Font(bold=True)

    widths = [4, 28, 28, 6, 6, 5, 7, 5, 5, 5, 6, 6, 7, 7, 8]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = w

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(out_path)
    print(f"Done! Saved {len(rows)} rows → {out_path}")

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("Usage: python3 parse_html_to_excel.py <saved_page.html>")
        sys.exit(1)
    parse(sys.argv[1])
